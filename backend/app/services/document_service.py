import os
import json
import fitz # PyMuPDF
from docx import Document as DocxDocument
from pptx import Presentation
from sqlalchemy.orm import Session
from langchain_text_splitters import RecursiveCharacterTextSplitter
from app.models.document import Document, DocumentChunk
from app.ai.embeddings import get_embeddings
from app.core.config import settings

class DocumentService:
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> list:
        # Returns list of tuples: (page_num, text)
        pages_content = []
        doc = fitz.open(file_path)
        ocr_engine = None
        for page_idx, page in enumerate(doc):
            text = page.get_text()
            # If native text extraction has very little content, use OCR (for scanned notes or slide decks)
            if not text or len(text.strip()) < 20:
                try:
                    if ocr_engine is None:
                        from rapidocr_onnxruntime import RapidOCR
                        ocr_engine = RapidOCR()
                    pix = page.get_pixmap()
                    img_bytes = pix.tobytes()
                    ocr_res, _ = ocr_engine(img_bytes)
                    if ocr_res:
                        ocr_lines = [line[1] for line in ocr_res if line and len(line) > 1]
                        if ocr_lines:
                            text = "\n".join(ocr_lines)
                except Exception:
                    pass

            if text and text.strip():
                pages_content.append((page_idx + 1, text.strip()))
        return pages_content

    @staticmethod
    def extract_text_from_docx(file_path: str) -> list:
        doc = DocxDocument(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        text = "\n".join(full_text)
        return [(1, text)] # DOCX usually processed as single stream or virtual page

    @staticmethod
    def extract_text_from_pptx(file_path: str) -> list:
        prs = Presentation(file_path)
        slides_content = []
        for slide_idx, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)
            text = "\n".join(slide_text)
            if text.strip():
                slides_content.append((slide_idx + 1, text))
        return slides_content

    @staticmethod
    def extract_text_from_txt(file_path: str) -> list:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        return [(1, text)]

    @staticmethod
    def process_and_store_document(db: Session, file_path: str, filename: str) -> Document:
        file_ext = os.path.splitext(filename)[1].lower().replace(".", "")
        if file_ext == "pdf":
            pages = DocumentService.extract_text_from_pdf(file_path)
        elif file_ext == "docx":
            pages = DocumentService.extract_text_from_docx(file_path)
        elif file_ext == "pptx":
            pages = DocumentService.extract_text_from_pptx(file_path)
        else: # Default txt
            pages = DocumentService.extract_text_from_txt(file_path)

        # Create Document record
        db_doc = Document(
            filename=filename,
            file_path=file_path,
            file_type=file_ext
        )
        db.add(db_doc)
        db.commit()
        db.refresh(db_doc)

        # Setup splitter
        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP
        )

        embeddings_service = get_embeddings()

        # Chunk and embed pages
        for page_num, text in pages:
            chunks = text_splitter.split_text(text)
            if not chunks:
                continue
            
            # Embed all chunks for this page
            try:
                vectors = embeddings_service.embed_documents(chunks)
            except Exception:
                # Fallback mock embedding dimension (384 for bge-small)
                vectors = [[0.0] * 384 for _ in chunks]

            for chunk_text, vector in zip(chunks, vectors):
                db_chunk = DocumentChunk(
                    document_id=db_doc.id,
                    page_number=page_num,
                    content=chunk_text,
                    embedding=json.dumps(vector)
                )
                db.add(db_chunk)
        
        db.commit()
        return db_doc

    @staticmethod
    def get_document(db: Session, doc_id: int) -> Document:
        return db.query(Document).filter(Document.id == doc_id).first()
